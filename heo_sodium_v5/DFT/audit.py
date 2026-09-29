"""audit.py - STEP 5: derived quantities + verdicts (spec section 7), and the pilot
checklist (spec section 4) via --pilot.

Frozen conventions (CLAUDE.md section 3, spec section 8):
    dE = (E_P3 - E_O3) * 1000 / 27  meV/f.u.,  dE > 0 = O3 wins
    V_avg = [E_desod + 13*E_Na - E_pris] / 13  > 0,  DFT E_Na ONLY (spec 8-5)
Scoring definitions are never touched here - findings are reported only (spec 8-7).

Usage:
    python audit.py --root $RUN_ROOT [--orb orb_on_dft.extxyz]     main audit
    python audit.py --root $RUN_ROOT --pilot                        section-4 checklist
"""
import argparse, glob, json, os, sys

import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import dft_common as C

XTAGS = ["x100", "x081", "x067", "x052"]
STEM = C.STEM_OF_XTAG

# |m| (mu_B) -> oxidation state bands; upper edges, first match wins.  Start values from
# the top-5 Tier-1 harvest (collect_dft.py) + spec section 7; FINALIZED AT PILOT.
MAG_BANDS = {
    "Ni": [(0.45, 4), (1.30, 3), (9.9, 2)],
    "Mn": [(3.40, 4), (9.9, 3)],
    "Co": [(0.70, 3), (2.00, None), (9.9, 2)],     # None = ambiguous, flagged not counted
    "V":  [(0.50, 5), (1.50, 4), (9.9, 3)],
    "Fe": [(9.9, 3)],
    "Cu": [(9.9, 2)],
}
SOFT_SLOPE_UNIFORM_STD = 0.05      # per-comp c spread below this -> "slope-type" (uniform)
SOFT_R2_MIN = 0.95


# ============================================================ loading
def load(root):
    man = pd.read_csv(os.path.join(root, "manifest.csv"), dtype=str).fillna("")
    raw = pd.read_csv(os.path.join(root, "dft_raw.csv"), dtype=str).fillna("")
    for c in ("E_dft", "E_relax", "p_order_dft", "max_force", "mlip_E", "mlip_dE"):
        raw[c] = pd.to_numeric(raw[c], errors="coerce")
    return man, raw


def e_na_dft(raw):
    r = raw[raw["set"] == "E_Na"]
    assert len(r), "E_Na row not collected yet - V audits impossible (spec 8-5 forbids MLIP E_Na)"
    e = float(r.iloc[0]["E_dft"]) / float(r.iloc[0]["natoms"])
    assert -1.6 < e < -1.0, f"DFT E_Na {e:.3f} eV/atom far from PBE ~-1.31 - wrong currency?"
    return e


def grid(raw):
    """{comp_id: {x_tag: {phase: raw-row}}} for the main sets (pilot excluded)."""
    g = {}
    for _, r in raw[raw["set"].isin(["top30", "anchor"])].iterrows():
        g.setdefault(r["comp_id"], {}).setdefault(r["x_tag"], {})[r["phase"]] = r
    return g


# ============================================================ audits (spec section 7 table)
def audit_convergence(g):
    rework = []
    for cid, xs in g.items():
        for xt, phases in xs.items():
            for ph, r in phases.items():
                if not (str(r["converged_relax"]) == "True" and str(r["converged_static"]) == "True"):
                    rework.append(f"{r['row_id']}: not converged")
                if r["phase_dft"] not in ("", ph):
                    rework.append(f"{r['row_id']}: phase flip {ph} -> {r['phase_dft']} "
                                  f"(P_order {r['p_order_dft']:+.2f})")
    return rework


def audit_sign(g):
    rows = []
    for cid, xs in g.items():
        rec = {"comp_id": cid}
        for xt in XTAGS:
            phases = xs.get(xt, {})
            if "O3" in phases and "P3" in phases:
                de = C.de_mev_fu(float(phases["O3"]["E_dft"]), float(phases["P3"]["E_dft"]))
                dm = float(phases["O3"]["mlip_dE"])
                rec[f"dE_dft_{xt}"] = de
                rec[f"dE_mlip_{xt}"] = dm
                rec[f"sign_ok_{xt}"] = bool(np.sign(de) == np.sign(dm))
        rows.append(rec)
    return pd.DataFrame(rows).set_index("comp_id")


def audit_voltage(g, e_na):
    rows = []
    for cid, xs in g.items():
        E_nat = {}
        for xt in XTAGS:
            phases = xs.get(xt, {})
            if phases:
                E_nat[STEM[xt]] = min(float(r["E_dft"]) for r in phases.values())
        segs = C.v_segments(E_nat, e_na)                  # unrounded for the identity check
        rec = {"comp_id": cid, **{k: round(v, 4) for k, v in segs.items()}}
        if "pris" in E_nat and "desod" in E_nat:
            va = C.v_avg(E_nat["pris"], E_nat["desod"], e_na)
            rec["V_avg_dft"] = round(va, 4)
            svals = [segs[f"V_seg_{C.X_LABEL[a]}_{C.X_LABEL[b]}"] for a, b in C.V_SEGMENTS]
            dns = [C.X_STEMS[b] - C.X_STEMS[a] for a, b in C.V_SEGMENTS]
            if all(np.isfinite(s) for s in svals):
                recon = sum(d * s for d, s in zip(dns, svals)) / 13.0
                rec["V_identity_resid"] = round(recon - va, 9)      # must be ~0 (spec section 7)
        rows.append(rec)
    return pd.DataFrame(rows).set_index("comp_id")


def audit_volume(g, man):
    """dV grammar per phase from DFT lattices, vs MLIP cells carried in the manifest."""
    mcell = {(r["row_id"]): r["mlip_cell"] for _, r in man.iterrows() if r["mlip_cell"]}
    rows = []
    for cid, xs in g.items():
        rec = {"comp_id": cid}
        for ph in ("O3", "P3"):
            rp, rd = xs.get("x100", {}).get(ph), xs.get("x052", {}).get(ph)
            if rp is None or rd is None or not rp["lattice"] or not rd["lattice"]:
                continue
            d = C.dV_report(json.loads(rp["lattice"]), json.loads(rd["lattice"]))
            for k in ("dV_pct", "dA_pct", "dh_perp_pct", "vm_strain"):
                rec[f"{k}_{ph}_dft"] = round(d[k], 4)
            cp, cd = mcell.get(rp["row_id"]), mcell.get(rd["row_id"])
            if cp and cd:
                dm = C.dV_report(json.loads(cp), json.loads(cd))
                for k in ("dV_pct", "dA_pct", "dh_perp_pct", "vm_strain"):
                    rec[f"{k}_{ph}_mlip"] = round(dm[k], 4)
        rows.append(rec)
    return pd.DataFrame(rows).set_index("comp_id")


def moments_of(r):
    """dft_raw mag_sites 'El:m;El:m;...' -> [(element, |m|)]."""
    out = []
    for tok in str(r["mag_sites"]).split(";"):
        if ":" in tok:
            el, m = tok.rsplit(":", 1)
            try:
                out.append((el, abs(float(m))))
            except ValueError:
                pass
    return out


def audit_magnetization(g, man):
    """Site moments at x=1 -> oxidation-state counts vs the stage-1 accounting (falsification
    test, spec question 2).  Thresholds in MAG_BANDS, finalized at pilot."""
    acct_of = {r["row_id"]: r["acct_json"] for _, r in man.iterrows()}
    rows = []
    for cid, xs in g.items():
        r = xs.get("x100", {}).get("O3")
        if r is None or not r["mag_sites"]:
            continue
        counts, amb = {}, 0
        for el, m in moments_of(r):
            if el not in MAG_BANDS:
                continue
            for hi, st in MAG_BANDS[el]:
                if m < hi:
                    if st is None: amb += 1
                    else: counts[(el, st)] = counts.get((el, st), 0) + 1
                    break
        acct = json.loads(acct_of.get(r["row_id"], "{}") or "{}")
        rows.append({
            "comp_id": cid,
            "n_Ni3_dft": counts.get(("Ni", 3), 0), "n_Ni3_mlip": acct.get("n_Ni3", ""),
            "n_Mn3_dft": counts.get(("Mn", 3), 0), "n_Mn3_mlip": acct.get("n_Mn3", ""),
            "n_Co2_dft": counts.get(("Co", 2), 0), "n_Co2_mlip": acct.get("n_Co2", ""),
            "n_V3_dft": counts.get(("V", 3), 0), "n_V3_mlip": acct.get("n_V3", ""),
            "n_V5_dft": counts.get(("V", 5), 0), "n_V5_mlip": acct.get("n_V5", ""),
            "n_ambiguous": amb,
        })
    return pd.DataFrame(rows).set_index("comp_id") if rows else pd.DataFrame()


def audit_softening(orb_path):
    """F0: force-component regression F_orb = c * F_dft on fixed geometries.
    Returns (global stats dict, per-comp DataFrame)."""
    from ase.io import read
    frames = read(orb_path, index=":")
    per_comp = {}
    for at in frames:
        cid = at.info.get("comp_id", "?")
        fd = at.get_array("dft_forces").ravel()
        fo = at.get_array("orb_forces").ravel()
        per_comp.setdefault(cid, [[], []])
        per_comp[cid][0].append(fd); per_comp[cid][1].append(fo)

    def fit(fd, fo):
        c = float(np.dot(fd, fo) / np.dot(fd, fd))          # least squares through origin
        resid = fo - c * fd
        ss = 1.0 - float(np.sum(resid**2) / np.sum((fo - fo.mean())**2))
        return c, float(np.mean(np.abs(resid))), ss

    rows = []
    all_fd, all_fo = [], []
    for cid, (fds, fos) in per_comp.items():
        fd, fo = np.concatenate(fds), np.concatenate(fos)
        all_fd.append(fd); all_fo.append(fo)
        c, cmae, r2 = fit(fd, fo)
        rows.append({"comp_id": cid, "c": round(c, 4), "cMAE": round(cmae, 4),
                     "R2": round(r2, 4), "n_forces": len(fd)})
    fd, fo = np.concatenate(all_fd), np.concatenate(all_fo)
    c, cmae, r2 = fit(fd, fo)
    df = pd.DataFrame(rows).set_index("comp_id")
    spread = float(df["c"].std())
    glob = {"c_global": round(c, 4), "cMAE_global": round(cmae, 4), "R2_global": round(r2, 4),
            "c_comp_std": round(spread, 4),
            "type": ("slope" if spread < SOFT_SLOPE_UNIFORM_STD and r2 > SOFT_R2_MIN else "scatter")}
    return glob, df


def audit_anchor_gate2(sign_df):
    """DFT edition of GATE 2: HOST vs HEO ordering at x067 (Zhao narrative: HEO delays O3->P3,
    so dE_x067(HEO) should exceed dE_x067(HOST)).  v1-only anchors carry no variant sigma, so
    the MLIP delta floor (5 meV/f.u.) is the significance threshold - recorded as a limitation."""
    out = {}
    try:
        host, heo = sign_df.loc["HOST_v1"], sign_df.loc["HEO_v1"]
        for cur in ("dft", "mlip"):
            d = heo[f"dE_{cur}_x067"] - host[f"dE_{cur}_x067"]
            out[f"gate2_diff_{cur}"] = round(float(d), 2)
            out[f"gate2_pass_{cur}"] = bool(d > C.DELTA_FLOOR_MEV_FU)
    except KeyError:
        out["gate2_note"] = "anchor x067 rows not complete yet"
    return out


# ============================================================ report + figures
def make_pptx(root, sign_df, volt_df, mag_df, soft, soft_df):
    """Figures as an editable .pptx with NATIVE charts (CLAUDE.md: never PNG alone)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.chart.data import XyChartData, CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs = Presentation()
    blank = prs.slide_layouts[6]

    def slide(title):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(9.4), Inches(0.5))
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].font.size = Pt(20)
        return s

    # 1. sign parity: mlip dE vs dft dE, one series per x
    s = slide("dE parity: MLIP vs DFT (meV/f.u.; quadrant = sign agreement)")
    cd = XyChartData()
    for xt in XTAGS:
        ser = cd.add_series(xt)
        for cid in sign_df.index:
            xm, yd = sign_df.loc[cid].get(f"dE_mlip_{xt}"), sign_df.loc[cid].get(f"dE_dft_{xt}")
            if pd.notna(xm) and pd.notna(yd):
                ser.add_data_point(float(xm), float(yd))
    s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(0.5), Inches(0.7),
                       Inches(9), Inches(6.3), cd)

    # 2. moment histogram (Ni, Mn) from the accounting audit
    if len(mag_df):
        s = slide("Charge accounting: MLIP prediction vs DFT site moments (x=1)")
        cd = CategoryChartData()
        cd.categories = list(mag_df.index)
        for col in ("n_Ni3_mlip", "n_Ni3_dft", "n_Mn3_mlip", "n_Mn3_dft"):
            cd.add_series(col, [float(v) if str(v) != "" else 0.0 for v in mag_df[col]])
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(0.7),
                           Inches(9), Inches(6.3), cd)

    # 3. softening per-composition slopes
    if soft_df is not None and len(soft_df):
        s = slide(f"F0 softening: c_global={soft['c_global']} ({soft['type']}-type, "
                  f"R2={soft['R2_global']}, comp spread {soft['c_comp_std']})")
        cd = CategoryChartData()
        cd.categories = list(soft_df.index)
        cd.add_series("c per composition", [float(v) for v in soft_df["c"]])
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(0.7),
                           Inches(9), Inches(6.3), cd)

    # 4. anchor V(x) staircase
    s = slide("Anchor voltage staircase V_seg(x) - DFT E_Na currency")
    cd = CategoryChartData()
    cd.categories = ["100-081", "081-067", "067-052"]
    for cid in ("HOST_v1", "HEO_v1"):
        if cid in volt_df.index:
            cd.add_series(cid, [float(volt_df.loc[cid].get(f"V_seg_{a}_{b}", np.nan))
                                for a, b in (("100", "081"), ("081", "067"), ("067", "052"))])
    s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(0.7),
                       Inches(9), Inches(6.3), cd)

    fp = os.path.join(root, "figs_dft.pptx")
    prs.save(fp)
    return fp


def main_audit(root, orb_path):
    man, raw = load(root)
    g = grid(raw)
    n_expected = len(man[man["set"].isin(["top30", "anchor"])])
    n_have = sum(len(p) for xs in g.values() for p in xs.values())
    print(f"audit over {n_have}/{n_expected} grid calcs, {len(g)} compositions")

    rework = audit_convergence(g)
    sign_df = audit_sign(g)
    e_na = e_na_dft(raw)
    volt_df = audit_voltage(g, e_na)
    vol_df = audit_volume(g, man)
    mag_df = audit_magnetization(g, man)
    soft, soft_df = (None, None)
    if orb_path and os.path.exists(orb_path):
        soft, soft_df = audit_softening(orb_path)
    gate2 = audit_anchor_gate2(sign_df)

    # ---- verdict (spec section 7 rule)
    top30_ids = set(man[man["set"] == "top30"]["comp_id"])
    s052 = sign_df.loc[[c for c in sign_df.index if c in top30_ids]]
    n_ok = int(s052.get("sign_ok_x052", pd.Series(dtype=bool)).fillna(False).sum())
    n_tot = int(s052.get("sign_ok_x052", pd.Series(dtype=bool)).notna().sum())
    slope_uniform = bool(soft and soft["type"] == "slope")
    finetune = not (n_ok >= 27 and slope_uniform)

    # ---- dft_verification.csv: one row per composition, {qty}_mlip/{qty}_dft pairs
    ver = sign_df.join([volt_df, vol_df] + ([mag_df] if len(mag_df) else []) +
                       ([soft_df.add_prefix("soft_")] if soft_df is not None else []), how="outer")
    ver.to_csv(os.path.join(root, "dft_verification.csv"))

    # ---- report
    lines = [
        "# report_dft.md - DFT verification (spec section 7)", "",
        f"- calcs audited: {n_have}/{n_expected}; compositions: {len(g)}",
        f"- DFT E_Na: {e_na:.4f} eV/atom (bcc, same currency; MLIP E_Na never used - spec 8-5)", "",
        "## Gates",
        f"- sign agreement x052 (top-30): **{n_ok}/{n_tot}** (gate: >= 27)",
        f"- softening: " + (f"c={soft['c_global']}, cMAE={soft['cMAE_global']}, "
                            f"R2={soft['R2_global']}, per-comp spread={soft['c_comp_std']} "
                            f"-> **{soft['type']}-type**" if soft else "orb_on_dft.extxyz not supplied"),
        f"- anchor GATE-2 (DFT): {gate2}",
        f"  (limitation: v1-only anchors -> no variant sigma; threshold = MLIP delta floor "
        f"{C.DELTA_FLOOR_MEV_FU} meV/f.u.)", "",
        "## Verdict (spec section 7 rule)",
        ("**Fine-tuning F1 NOT triggered** - ranking confirmed; c recorded."
         if not finetune else
         "**Fine-tuning F1 TRIGGERED** - replace the STEP-1 list with the training set, "
         "reuse STEPs 2-4; anchors permanently excluded from training."), "",
        "## Rework queue (unconverged / phase-flipped)",
    ]
    lines += [f"- {r}" for r in rework] or ["- none"]
    lines += ["", "## Voltage sanity",
              f"- all V_seg > 0: {bool((volt_df.filter(like='V_seg_').fillna(1) > 0).values.all())}",
              f"- identity resid max |sum(dn*V_seg)/13 - V_avg|: "
              f"{volt_df.get('V_identity_resid', pd.Series(dtype=float)).abs().max()}"]
    if "HOST_v1" in volt_df.index:
        va = volt_df.loc["HOST_v1"].get("V_avg_dft", np.nan)
        lines.append(f"- HOST V_avg = {va} V (window 2.5-3.4, expect ~3.1)")
    per_x = {xt: f"{int(sign_df.get(f'sign_ok_{xt}', pd.Series(dtype=bool)).fillna(False).sum())}"
                 f"/{int(sign_df.get(f'sign_ok_{xt}', pd.Series(dtype=bool)).notna().sum())}"
             for xt in XTAGS}
    lines += ["", "## Sign agreement per x", f"- {per_x}", "",
              "## Notes",
              "- phase read-out uses prismatic_order (heo_worker's relaxation-robust classifier); "
              "the letter classifier asserts on buckled cells (deviation from spec 3-3 wording).",
              "- representative vacancy orderings (k_best) were chosen in the MLIP currency "
              "(spec section 1 limitation paragraph)."]
    C.atomic_write(os.path.join(root, "report_dft.md"), "\n".join(lines) + "\n")

    fp = make_pptx(root, sign_df, volt_df, mag_df, soft or {"c_global": "n/a", "type": "n/a",
                                                            "R2_global": "n/a", "c_comp_std": "n/a"},
                   soft_df)
    print(f"wrote dft_verification.csv, report_dft.md, {os.path.basename(fp)}")
    print(f"sign x052: {n_ok}/{n_tot}; fine-tune triggered: {finetune}")


# ============================================================ pilot checklist (spec section 4)
def pilot_audit(root):
    man, raw = load(root)
    p = raw[raw["set"] == "pilot"].copy()
    print(f"pilot rows collected: {len(p)}/8")
    checks = []

    def base_id(rid):
        return rid.rsplit("__", 1)[0]

    # 1. gam vs std: |d(dE)| < 5 meV/f.u. at both x, AND force MAE < 0.02 eV/A on relax step 0
    for xt in ("x100", "x052"):
        des = {}
        for mode in ("std", "gam"):
            sel = p[(p["x_tag"] == xt) & (p["vasp_mode"] == mode)]
            o3 = sel[sel["phase"] == "O3"]; p3 = sel[sel["phase"] == "P3"]
            if len(o3) and len(p3):
                des[mode] = C.de_mev_fu(float(o3.iloc[0]["E_dft"]), float(p3.iloc[0]["E_dft"]))
        if len(des) == 2:
            d = abs(des["std"] - des["gam"])
            checks.append((f"gam vs std |d(dE)| {xt} = {d:.2f} meV/f.u.", d < 5.0))
        else:
            checks.append((f"gam vs std dE {xt}: rows incomplete", False))

    maes = []
    for _, r_std in p[p["vasp_mode"] == "std"].iterrows():
        twin = p[(p["vasp_mode"] == "gam") & (p["row_id"] == base_id(r_std["row_id"]) + "__gam")]
        if not len(twin):
            continue
        try:
            f0 = []
            for rid in (r_std["row_id"], twin.iloc[0]["row_id"]):
                d = json.load(open(os.path.join(root, "raw", f"{rid}.json")))
                f0.append(np.asarray(
                    d["relax"]["calcs_reversed"][0]["output"]["ionic_steps"][0]["forces"], float))
            maes.append(float(np.mean(np.abs(f0[0] - f0[1]))))
        except Exception as e:
            checks.append((f"force MAE {base_id(r_std['row_id'])}: unreadable ({e})", False))
    if maes:
        checks.append((f"gam vs std force MAE (relax step 0, identical geometry) = "
                       f"{max(maes):.4f} eV/A", max(maes) < 0.02))

    # 2. MAGMOM convergence vs accounting (x=1 std rows)
    acct_of = {r["row_id"]: r["acct_json"] for _, r in man.iterrows()}
    for _, r in p[(p["x_tag"] == "x100") & (p["vasp_mode"] == "std")].iterrows():
        acct = json.loads(acct_of.get(r["row_id"], "{}") or "{}")
        ms = [m for el, m in moments_of(r) if el == "Ni"]
        n3 = sum(1 for m in ms if 0.45 <= m < 1.3)
        checks.append((f"{r['row_id']}: n_Ni3(moments)={n3} vs accounting={acct.get('n_Ni3')} "
                       f"(thresholds are pilot-tunable)", True))

    # 3. custodian recovery on the NELM=8 test row
    test = man[man["note"].str.contains("custodian_test", na=False)]
    if len(test):
        rid = test.iloc[0]["row_id"]
        cj = glob.glob(os.path.join(root, "calcs", rid, "run-*", "**", "custodian.json"),
                       recursive=True)
        fired = False
        for fp in cj:
            try:
                fired = fired or any(c.get("corrections") for j in json.load(open(fp))
                                     for c in [j] if isinstance(j, dict))
            except Exception:
                pass
        checks.append((f"custodian recovery fired on {rid} ({len(cj)} custodian.json found)", fired))

    # 4. TaskDoc field presence (spec 10-#3)
    miss = set()
    for _, r in p.iterrows():
        if r["missing_fields"]:
            miss |= set(str(r["missing_fields"]).split(","))
    checks.append((f"TaskDoc fields all present (missing: {sorted(miss) or 'none'})", not miss))

    # 5. wall time T -> suggested t_est_hours (claim started_at -> DONE finished_at)
    ts = []
    for _, r in p.iterrows():
        try:
            calc = os.path.join(root, "calcs", r["row_id"])
            t0 = json.load(open(os.path.join(calc, ".claim", "claim.json")))["started_at"]
            t1 = json.loads(open(os.path.join(calc, "DONE")).read())["finished_at"]
            ts.append((pd.Timestamp(t1) - pd.Timestamp(t0)).total_seconds() / 3600.0)
        except Exception:
            pass
    if ts:
        t = max(ts)
        checks.append((f"T per calc: max {t:.2f} h, mean {np.mean(ts):.2f} h -> set "
                       f"t_est_hours ~ {1.3 * t:.1f} in config.yaml; total ~ 256*T/5 = "
                       f"{256 * np.mean(ts) / 5:.0f} h wall", True))

    print("\n=== PILOT CHECKLIST (spec section 4) ===")
    ok = True
    for msg, passed in checks:
        print(f"  [{'PASS' if passed else 'FAIL'}] {msg}")
        ok &= bool(passed)
    print("\n=> " + ("all checks PASS - set kmode (std|gam) + t_est_hours in config.yaml, "
                     "then ./submit.sh for the main run" if ok else
                     "FAILURES above - fix before the main run (spec: no main run without pilot pass)"))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--orb", default=None, help="orb_on_dft.extxyz from the GPU machine")
    ap.add_argument("--pilot", action="store_true")
    args = ap.parse_args()
    root = os.path.abspath(args.root)
    if args.pilot:
        pilot_audit(root)
    else:
        main_audit(root, args.orb or os.path.join(root, "orb_on_dft.extxyz"))


if __name__ == "__main__":
    main()
